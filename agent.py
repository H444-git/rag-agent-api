import os

from dotenv import load_dotenv

load_dotenv()

import requests
import json
import numpy as np
import fitz
from text2vec import SentenceModel
from langchain_text_splitters import RecursiveCharacterTextSplitter
from tavily import TavilyClient
from logger_utils import logger
from docx import Document

# ===== 配置 =====
DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY")
TAVILY_API_KEY = os.getenv("TAVILY_API_KEY")
url = "https://api.deepseek.com/v1/chat/completions"

# ===== 初始化 Tavily 客户端 =====
tavily_client = TavilyClient(api_key=TAVILY_API_KEY)

SYSTEM_PROMPT = """
你是一个专业的技术顾问，也是一个自主决策的智能助手。

行为准则：
1. 请在回答用户问题之前，在 <thinking> 标签中简要分析用户的需求和你的决策思路，还要规划出需要调用哪些工具，按什么顺序调用。
2. 如果用户的问题涉及实时信息（如天气、新闻、最新事件），请调用搜索工具来获取信息。
3. 如果用户的问题基于通用知识或你已有信息，可以直接基于你的知识回答。
4. 如果用户的问题不清晰，你可以反问澄清，而不是直接猜测。
5. 回答要友好、清晰、简洁，风格自然。
6. 每次回答后，你可以询问用户是否满意或是否还有问题，但不是必须。
7. 当你认为任务已经完成时，请以【H,我可以给出答案啦！】开头
8.对于需要多个步骤才能完成的任务（例如：需要先搜索信息，再对信息进行计算），请在 <thinking> 标签中先列出完成该任务所需的步骤清单，然后按清单顺序依次执行。

示例：
用户问题：北京和上海温度数字之和
步骤清单：
1. 搜索北京实时温度
2. 搜索上海实时温度
3. 从搜索结果中提取温度数字
4. 计算两个数字之和
5. 输出结果
注意：你是一个智能助手，可以根据需要自主决定是否调用工具、是否反问，以及如何组织回答。
"""

tools = [
    {
        "type": "function",
        "function": {
            "name": "search_web",
            "description": "搜索互联网获取实时信息。当用户的问题涉及新闻、天气、最新事件、实时数据或其他需要联网才能获取的信息时，使用此工具。此工具可以被多次调用",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "用于搜索的精确关键词，例如'重庆天气'或'Python 2024 最新特性'。",
                    }
                },
                "required": ["query"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "calculate",
            "description": "执行数学计算。当用户需要计算数值表达式、百分比、单位换算或其他数学运算时使用。",
            "parameters": {
                "type": "object",
                "properties": {
                    "expression": {
                        "type": "string",
                        "description": "要计算的数学表达式，例如 '100 * 0.15' 或 '5000 / 12'",
                    }
                },
                "required": ["expression"],
            },
        },
    },
]
# ===== 加载模型 =====
print("正在加载模型...")
model = SentenceModel()
print("模型加载完成！")


def split_document(text, chunk_size=100, chunk_overlap=50):
    """文档分块函数"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""],
    )
    return splitter.split_text(text)


def load_knowledge(file_path, chunk_size=100):
    """加载知识库函数"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
        chunks = split_document(text, chunk_size)
        return [f"块{i+1} {chunk}" for i, chunk in enumerate(chunks)]
    except FileNotFoundError:
        print(f"文件 {file_path} 不存在")
        return []


def parse_pdf(file_bytes):
    """解析 PDF 字节流，返回纯文本"""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text()
    doc.close()
    return text


def parse_docx(file_bytes):
    """解析 Word 字节流，返回纯文本"""
    from io import BytesIO

    doc = Document(BytesIO(file_bytes))
    text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    logger.info(f"解析 docx，共 {len(doc.paragraphs)} 个段落，文本长度 {len(text)}")
    return text


def parse_file(filename, file_bytes):
    """
    根据文件扩展名选择解析方式，返回纯文本
    """
    if filename.endswith(".txt"):
        return file_bytes.decode("utf-8")
    elif filename.endswith(".pdf"):
        return parse_pdf(file_bytes)
    elif filename.endswith(".docx"):
        return parse_docx(file_bytes)
    else:
        raise ValueError(f"不支持的文件类型：{filename}")


knowledge = load_knowledge("知识库简易版.txt", chunk_size=150)
knowledge_vecs = model.encode(knowledge) if knowledge else []

import faiss

# 将知识向量转为 float32（FAISS 要求的数据类型）
knowledge_vecs_faiss = np.array(knowledge_vecs).astype("float32")

# 获取向量维度
dim = knowledge_vecs_faiss.shape[1]

# 创建 FAISS 索引（使用内积索引，配合归一化实现余弦相似度）
index = faiss.IndexFlatIP(dim)  # IP = Inner Product（内积）
# 归一化向量，使内积等于余弦相似度
faiss.normalize_L2(knowledge_vecs_faiss)
index.add(knowledge_vecs_faiss)

print(f"✅ FAISS 索引已创建，包含 {index.ntotal} 条知识向量")


def get_all_knowledge():
    """返回当前知识库里的所有文本块"""
    return knowledge


def add_knowledge(new_chunks):
    """
    向知识库追加新的文本块，并同步更新 FAISS 索引
    返回：新增的块数、当前总块数
    """
    global knowledge

    # 1. 向量化
    new_vecs = model.encode(new_chunks)
    new_vecs_faiss = np.array(new_vecs).astype("float32")

    # 2. 归一化（FAISS 要求）
    faiss.normalize_L2(new_vecs_faiss)

    # 3. 追加到 knowledge 和 index，保证顺序一致
    knowledge.extend(new_chunks)
    index.add(new_vecs_faiss)

    # 打印新增的每一块内容
    for i, chunk in enumerate(new_chunks):
        logger.info(f"新增块 {i+1}: {chunk}...")

    logger.info(f"知识库更新：新增 {len(new_chunks)} 块，当前共 {len(knowledge)} 块")

    return len(new_chunks), len(knowledge)


def search_mostk(query_vec, K=16, threshold=0.4):
    """
    使用 FAISS 检索最相关的 K 条知识
    """
    # 将查询向量转为 float32 并归一化
    query_vec_faiss = np.array([query_vec]).astype("float32")
    faiss.normalize_L2(query_vec_faiss)

    # FAISS 检索
    distances, indices = index.search(query_vec_faiss, K)

    # distances 是余弦相似度（因为向量已归一化）
    # indices 是知识库中的索引位置

    matched = []
    for i, (dist, idx) in enumerate(zip(distances[0], indices[0])):
        if idx == -1:
            continue
        sim = float(dist)  # dist 就是余弦相似度
        logger.debug(f"🔍 知识块 {idx+1} 的注意力得分: {sim:.4f}")
        if sim >= threshold:
            matched.append(knowledge[idx])
        else:
            logger.debug(f"    ⏭️  知识块 {idx+1} 低于阈值 {threshold}，已跳过")
        logger.info(f"检索完成:命中{len(matched)}条知识")
    return matched[:K]


def web_search_tavily(query):
    """
    使用 Tavily 进行联网搜索，返回结构化的搜索结果
    """
    try:
        response = tavily_client.search(
            query=query,
            search_depth="basic",  # basic 消耗 1 个信用点，advanced 消耗 2 个
            max_results=3,
            include_answer=True,
            include_raw_content=False,
            include_images=False,
        )

        results = response.get("results", [])
        answer = response.get("answer", "")

        if not results and not answer:
            return "未找到相关搜索结果"

        search_summary = []
        if answer:
            search_summary.append(f"📌 摘要：{answer}\n")

        for i, item in enumerate(results[:3], 1):
            title = item.get("title", "无标题")
            content = item.get("content", "")
            url = item.get("url", "")
            search_summary.append(
                f"{i}. {title}\n   {content[:200]}...\n   来源：{url}"
            )

        return "\n".join(search_summary)

    except Exception as e:
        return f"搜索失败: {str(e)}"


def calculate(expression):
    """
    安全地计算数学表达式
    """
    # 只允许数字、运算符、括号、空格
    allowed_chars = set("0123456789+-*/().% ")
    if not all(c in allowed_chars for c in expression):
        return "错误：表达式包含不支持的字符"

    try:
        # 使用 eval 计算，但限制全局和局部命名空间为空，避免执行危险代码
        result = eval(expression, {"__builtins__": {}}, {})
        return f"计算结果：{expression} = {result}"
    except Exception as e:
        return f"计算错误：{str(e)}"


def call_llm(prompt, mode, tools=None):
    """
    调用大模型
    mode:"chat"或"search"(提取关键词)
    tools:工具列表，如果提供则模型可以调用工具
    """
    if mode == "search":
        search_system = (
            "你是一个搜索关键词提取专家。只输出关键词，用空格分隔，不要输出其他内容。"
        )
        messages = [
            {"role": "system", "content": search_system},
            {"role": "user", "content": f"用户问题:{prompt}\n关键词:"},
        ]
        data = {
            "model": "deepseek-chat",
            "messages": messages,
        }
    else:
        messages = [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ]
        data = {
            "model": "deepseek-chat",
            "messages": messages,
        }
        # 如果有tools，添加到请求中
        if tools:
            data["tools"] = tools
            data["tool_choice"] = "auto"  # 模型自己决定是否调用工具

    use_stream = mode == "chat" and tools is None
    data["stream"] = use_stream

    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
    }
    response = requests.post(url, headers=headers, json=data, stream=use_stream)

    if use_stream:
        full_content = ""
        for line in response.iter_lines():
            if line:
                line = line.decode("utf-8")
                if line.startswith("data: "):
                    try:
                        chunk = json.loads(line[6:])
                        delta = chunk.get("choices", [{}])[0].get("delta", {})
                        content = delta.get("content", "")
                        if content:
                            full_content += content
                            print(content, end="", flush=True)  # 实时打印
                    except json.JSONDecodeError:
                        continue
        return full_content
    else:
        result = response.json()
        return result


def rag_chat(message, history):
    logger.info(f"收到问题:{message}")
    if not knowledge:
        logger.warning("知识库为空")
        return "知识库为空，请检查 知识库简易版.txt 文件"

    # 1. 先检索本地知识库
    query_vec = model.encode(message)

    # 如果问题里包含“几步”“哪些步骤”“流程”等词，用更大的 K
    if any(word in message for word in ["几步", "步骤", "流程", "哪些"]):
        K = 16
    else:
        K = 6

    matched = search_mostk(query_vec, K)

    context = {
        "original_question": message,
        "knowledge": "\n".join(matched) if matched else "未找到相关知识",
        "tool_results": [],
        "round_count": 0,
        "max_rounds": 3,
    }

    while context["round_count"] < context["max_rounds"]:
        tool_summary = (
            "\n".join(
                [
                    f"第{t['tool']}次搜索：{t['result'][:200]}..."
                    for t in context["tool_results"]
                ]
            )
            if context["tool_results"]
            else "暂无"
        )
        user_prompt = f"""
    【知识库内容】
    {context['knowledge']}

    【已有工具调用结果】
    {tool_summary}

    【用户问题】
    {message}

    请根据以上信息回答用户问题。如果知识库内容不足，且问题涉及实时信息，你可以调用 search_web 工具来获取更多信息。
    """

        result = call_llm(user_prompt, "chat", tools=tools)

        if isinstance(result, dict) and "error" in result:
            logger.error(f"模型调用失败:{result['error']}")
            return result["error"]

        # 检查是否是工具调用
        if isinstance(result, dict) and "choices" in result:
            choice = result["choices"][0]
            message_obj = choice["message"]

            if "tool_calls" in message_obj and message_obj["tool_calls"]:
                tool_call = message_obj["tool_calls"][0]
                tool_name = tool_call["function"]["name"]
                tool_args = json.loads(tool_call["function"]["arguments"])

                logger.info(f"🔧 模型请求调用工具: {tool_name}")
                logger.info(f"📝 工具参数: {tool_args}")

                if tool_name == "search_web":
                    search_results = web_search_tavily(tool_args["query"])
                    logger.info(f"搜索完成,返回{search_results}字符")
                    context["tool_results"].append(
                        {"tool": tool_name, "args": tool_args, "result": search_results}
                    )
                    context["round_count"] += 1
                    context["last_tool_call"] = {
                        "tool_name": tool_name,
                        "args": tool_args,
                    }
                    continue
                elif tool_name == "calculate":
                    result = calculate(tool_args["expression"])
                    logger.info(f"计算结果:{result}")
                    context["tool_results"].append(
                        {"tool": tool_name, "args": tool_args, "result": result}
                    )
                    context["round_count"] += 1
                    continue
                else:
                    logger.error(f"未知工具:{tool_name}")
                    return f"未知工具: {tool_name}"
            else:
                # 直接回答
                content = message_obj.get("content", "")
                logger.info("回答完成！")
                return content
        elif isinstance(result, str):
            # 流式返回的字符串
            logger.info("回答完成(流式)")
            return result
    logger.warning(f"达到最大轮次{context['max_rounds']},任务未完成")
    return f"经过{context['max_rounds']}轮尝试，任务未完成。已获取的信息：{context['tool_results']}"
